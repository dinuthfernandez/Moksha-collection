"""Add a "Kids Size" column to Zoho_Items_Export.xlsx.

Only rows whose Size column is "Kids" are filled; the age range (e.g. 0-1,
2-3) is read from the end of the SKU, e.g. "FRK-KASAV-BLUE-001- 0 - 1" -> "0-1".
Every other row stays blank.
"""
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_PATH = r"C:\Users\Asus\Desktop\Moksha collection\Zoho_Items_Export.xlsx"
FALLBACK_PATH = r"C:\Users\Asus\Desktop\Moksha collection\Zoho_Items_Export_KidsSize.xlsx"
VALID_KIDS_SIZES = {"0-1", "1-2", "2-3", "3-4", "4-5", "6-7"}


def kids_size_from_sku(sku: str) -> str:
    m = re.search(r"(\d{1,2})\s*-\s*(\d{1,2})\s*$", sku.strip())
    if not m:
        return ""
    value = f"{int(m.group(1))}-{int(m.group(2))}"
    return value if value in VALID_KIDS_SIZES else ""


def main() -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    sku_col = headers.index("SKU") + 1
    size_col = headers.index("Size") + 1

    if "Kids Size" in headers:
        kids_col = headers.index("Kids Size") + 1
    else:
        kids_col = len(headers) + 1
        cell = ws.cell(row=1, column=kids_col, value="Kids Size")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(kids_col)].width = 12

    filled = missing = 0
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=kids_col)
        if ws.cell(row=row_idx, column=size_col).value == "Kids":
            value = kids_size_from_sku(str(ws.cell(row=row_idx, column=sku_col).value or ""))
            cell.value = value or None
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if value:
                filled += 1
            else:
                missing += 1
                print("No kids size found for SKU:", ws.cell(row=row_idx, column=sku_col).value)
        else:
            cell.value = None

    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    try:
        wb.save(EXCEL_PATH)
        print("Saved", EXCEL_PATH)
    except PermissionError:
        wb.save(FALLBACK_PATH)
        print("Original is open in Excel; saved a copy to", FALLBACK_PATH)
    print(f"Kids Size filled: {filled}, kids rows without a size: {missing}")


if __name__ == "__main__":
    main()
