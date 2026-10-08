"""Add a Website Serial column to Zoho_Items_Export.xlsx.

Groups items that are the "same design" (same product/material/style,
differing only by color and/or size) and assigns each distinct design
group a unique serial number (zero-padded to 5 digits), matching Zoho's
live "Website serial" custom field (cf_website_serial, data_type=number).

Design key = Item Name (group) + SKU tokens with recognized color and
size tokens stripped out, in original order. Two SKUs land in the same
design group only if their remaining ("design") tokens are identical.
"""
import re
import collections

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EXCEL_PATH = r"C:\Users\Asus\Desktop\Moksha collection\Zoho_Items_Export.xlsx"
GUIDE_PATH = r"C:\Users\Asus\Desktop\Moksha collection\Colour code Guide.xlsx"

# ---- Color token recognition (same source of truth as fill_colors_from_sku.py) ----
EXTRA_COLOR_CODES = {
    "BLK", "BLU", "BLUE", "YLW", "YLE", "MX", "WT", "GLD", "GOLDEN", "SILVER",
    "NVY", "NAVY", "INDG", "OLV", "ORNG", "RPNK", "GREEN", "GREY", "GR", "PUP",
}
COMBO_COLOR_PAIRS = {
    ("DRK", "BLU"), ("DRK", "BLUE"), ("DRK", "GRN"), ("DRK", "GREEN"),
    ("LMN", "YLW"), ("MST", "YLW"),
}

# ---- Size token recognition (same source of truth as fill_sizes_from_sku.py) ----
ALPHA_SIZES = {"XS", "S", "M", "L", "XL", "XXL", "2XL", "XXXL", "3XL", "XXXXL", "4XL", "5XL"}
KIDS_PAIRS = {(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (6, 7)}


def load_color_codes() -> set:
    wb = openpyxl.load_workbook(GUIDE_PATH, data_only=True)
    ws = wb["Sheet1"]
    codes = set(EXTRA_COLOR_CODES)
    for row in ws.iter_rows(min_row=5, values_only=True):
        code = row[1]
        if code:
            codes.add(code.strip().upper())
    return codes


def tokens_of(sku: str):
    return [t for t in re.split(r"[^A-Za-z0-9]+", sku) if t]


def build_design_key(item_name: str, sku: str, color_codes: set) -> str:
    raw = (sku or "").strip()
    toks = tokens_of(raw)
    up_toks = [t.upper() for t in toks]
    n = len(up_toks)

    # Mark indices that are "kids age range" size pair at the very end (e.g. "3","4").
    drop = [False] * n
    if n >= 2:
        a_raw, b_raw = up_toks[-2], up_toks[-1]
        if a_raw.isdigit() and b_raw.isdigit():
            a, b = int(a_raw), int(b_raw)
            if (a, b) in KIDS_PAIRS:
                drop[-2] = True
                drop[-1] = True

    for i, tok in enumerate(up_toks):
        if drop[i]:
            continue
        if i + 1 < n and (tok, up_toks[i + 1]) in COMBO_COLOR_PAIRS:
            drop[i] = True
            drop[i + 1] = True
            continue
        if tok in color_codes:
            drop[i] = True
            continue
        if tok in ALPHA_SIZES:
            drop[i] = True
            continue
        # digit-prefixed concatenated size like "001L" or "0001M"
        m = re.match(r"^\d+(XS|S|M|L|XL|XXL|2XL|XXXL|3XL|XXXXL|4XL|5XL)$", tok)
        if m:
            drop[i] = True

    remaining = [tok for i, tok in enumerate(up_toks) if not drop[i]]
    return (item_name or "").strip().upper() + "||" + "-".join(remaining)


def main() -> None:
    color_codes = load_color_codes()

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    headers = [c.value for c in ws[1]]

    name_col = headers.index("Item Name") + 1
    sku_col = headers.index("SKU") + 1

    if "Website Serial" in headers:
        serial_col = headers.index("Website Serial") + 1
    else:
        serial_col = len(headers) + 1
        header_cell = ws.cell(row=1, column=serial_col, value="Website Serial")
        header_cell.font = Font(bold=True, color="FFFFFF")
        header_cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(serial_col)].width = 16

    design_key_to_serial = {}
    next_serial = 1
    group_members = collections.defaultdict(list)

    row_keys = []
    for row_idx in range(2, ws.max_row + 1):
        item_name = ws.cell(row=row_idx, column=name_col).value
        sku = ws.cell(row=row_idx, column=sku_col).value
        key = build_design_key(item_name, sku, color_codes)
        row_keys.append((row_idx, key, sku))
        group_members[key].append(sku)

    for row_idx, key, sku in row_keys:
        if key not in design_key_to_serial:
            design_key_to_serial[key] = next_serial
            next_serial += 1
        serial = design_key_to_serial[key]
        cell = ws.cell(row=row_idx, column=serial_col, value=f"{serial:05d}")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    wb.save(EXCEL_PATH)

    total_groups = len(design_key_to_serial)
    multi_groups = sum(1 for members in group_members.values() if len(members) > 1)
    solo_groups = total_groups - multi_groups
    largest = max(group_members.values(), key=len)

    print(f"Total rows: {len(row_keys)}")
    print(f"Total distinct design groups (unique Website Serial numbers): {total_groups}")
    print(f"Groups with multiple color/size variants: {multi_groups}")
    print(f"Groups with only a single item (unique design, no siblings): {solo_groups}")
    print(f"Largest design group size: {len(largest)} -> sample SKUs: {largest[:6]}")


if __name__ == "__main__":
    main()
