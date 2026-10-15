"""Push ONLY the "Kids Size" column from Zoho_Items_Export.xlsx to the Zoho
"Kids size" custom field (cf_kids_size). No other item attribute is touched."""
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from app.config import get_settings
from app.services.zoho_inventory import ZohoInventoryClient

EXCEL_PATH = r"C:\Users\Asus\Desktop\Moksha collection\Zoho_Items_Export.xlsx"
FALLBACK_PATH = r"C:\Users\Asus\Desktop\Moksha collection\Zoho_Items_Export_KidsSize.xlsx"
KIDS_SIZE_FIELD_ID = "6877211000005984001"


def main() -> None:
    settings = get_settings()
    client = ZohoInventoryClient(
        client_id=settings.zoho_client_id,
        client_secret=settings.zoho_client_secret,
        refresh_token=settings.zoho_refresh_token,
        organization_id=settings.zoho_organization_id,
    )
    sku_to_id = {
        str(it["sku"]).strip(): it["item_id"]
        for it in client.fetch_all_items()
        if it.get("sku") and it.get("item_id")
    }

    path = EXCEL_PATH
    if "Kids Size" not in [c.value for c in openpyxl.load_workbook(path, read_only=True).active[1]]:
        path = FALLBACK_PATH
    ws = openpyxl.load_workbook(path, data_only=True).active
    headers = [c.value for c in ws[1]]
    sku_i, kids_i = headers.index("SKU"), headers.index("Kids Size")

    pushed, failed, not_found = 0, {}, []
    for row in ws.iter_rows(min_row=2, values_only=True):
        kids, sku = row[kids_i], row[sku_i]
        if not kids or not sku:
            continue
        sku = str(sku).strip()
        item_id = sku_to_id.get(sku)
        if not item_id:
            not_found.append(sku)
            continue
        body = {"custom_fields": [{"customfield_id": KIDS_SIZE_FIELD_ID, "value": str(kids)}]}
        for attempt in range(1, 4):
            try:
                client._request_json(
                    f"/items/{item_id}",
                    {"organization_id": settings.zoho_organization_id},
                    method="PUT",
                    json_body=body,
                )
                pushed += 1
                break
            except Exception as exc:  # noqa: BLE001
                if attempt == 3:
                    failed[sku] = str(exc)[:200]
                else:
                    time.sleep(2 * attempt)
        time.sleep(0.35)

    print(f"Pushed: {pushed}, not found: {len(not_found)}, failed: {len(failed)}")
    for sku in not_found:
        print("  not found:", sku)
    for sku, err in failed.items():
        print("  failed:", sku, err)


if __name__ == "__main__":
    main()
