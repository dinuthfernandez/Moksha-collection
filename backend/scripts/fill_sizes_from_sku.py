"""Add a Size column to Zoho_Items_Export.xlsx.

Fills size for Clothing-category rows only, derived from SKU patterns,
mapped to the exact live Zoho "Size" custom-field dropdown values:
XS, S, M, L, XL, 2XL, 3XL, 4XL, 5XL, Kids (plus Ring/Necklace/Bracelet/
Bangle values which are not applicable to clothing).

Accessories rows are left blank (size not applicable).
"""
import re

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EXCEL_PATH = r"C:\Users\Asus\Desktop\Moksha collection\Zoho_Items_Export.xlsx"

ALPHA_SIZES = ["XXXXL", "XXXL", "XXL", "XL", "5XL", "4XL", "3XL", "2XL", "XS", "S", "M", "L"]
ALPHA_CANON = {
    "XS": "XS", "S": "S", "M": "M", "L": "L", "XL": "XL",
    "XXL": "2XL", "2XL": "2XL",
    "XXXL": "3XL", "3XL": "3XL",
    "XXXXL": "4XL", "4XL": "4XL",
    "5XL": "5XL",
}
KIDS_PAIRS = {(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (6, 7)}

VALID_SIZES = {"XS", "S", "M", "L", "XL", "2XL", "3XL", "4XL", "5XL", "Kids"}


def detect_size(sku: str) -> str:
    raw = sku.strip()
    # Kids age-range check at the very end, e.g. "...-3-4", "...- 0 - 1"
    m = re.search(r"(\d{1,2})\s*-\s*(\d{1,2})$", raw)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        if (a, b) in KIDS_PAIRS:
            return "Kids"
    tokens = [t for t in re.split(r"[^A-Za-z0-9]+", raw) if t]
    for tok in reversed(tokens):
        up = tok.upper()
        if up in ALPHA_CANON:
            return ALPHA_CANON[up]
        m2 = re.match(r"^\d*(" + "|".join(ALPHA_SIZES) + r")$", up)
        if m2:
            return ALPHA_CANON[m2.group(1)]
    return ""


def main() -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    headers = [c.value for c in ws[1]]

    sku_col = headers.index("SKU") + 1
    cat_col = headers.index("Product Category") + 1

    if "Size" in headers:
        size_col = headers.index("Size") + 1
    else:
        size_col = len(headers) + 1
        header_cell = ws.cell(row=1, column=size_col, value="Size")
        ref_header = ws.cell(row=1, column=1)
        header_cell.font = Font(bold=True, color="FFFFFF")
        header_cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(size_col)].width = 10

    filled = 0
    blank_clothing = 0
    accessories = 0
    invalid = 0

    for row_idx in range(2, ws.max_row + 1):
        category = ws.cell(row=row_idx, column=cat_col).value
        sku = ws.cell(row=row_idx, column=sku_col).value or ""
        cell = ws.cell(row=row_idx, column=size_col)
        if category == "Clothing":
            size = detect_size(sku)
            if size and size in VALID_SIZES:
                cell.value = size
                cell.alignment = Alignment(horizontal="center", vertical="center")
                filled += 1
            else:
                cell.value = None
                blank_clothing += 1
                if size and size not in VALID_SIZES:
                    invalid += 1
        else:
            cell.value = None
            accessories += 1

    wb.save(EXCEL_PATH)

    print(f"Clothing rows with size filled: {filled}")
    print(f"Clothing rows left blank (no detectable size): {blank_clothing}")
    print(f"Accessories rows (size left blank): {accessories}")
    print(f"Invalid detections skipped: {invalid}")


if __name__ == "__main__":
    main()
