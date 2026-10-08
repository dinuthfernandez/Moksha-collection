"""
Exports the full Zoho Inventory items list to an Excel file with:
Item Name, SKU, Product Category, Product Sub Category, Show in Online Store.

Usage:
    .venv\\Scripts\\python.exe scripts\\export_zoho_items_excel.py [output_path.xlsx]

Notes:
- The Zoho Inventory /items list endpoint does not return custom fields
  (Product Category / Product Sub Category), so each item's detail is
  fetched individually (concurrently) to read its custom field values.
"""
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.services.zoho_inventory import ZohoInventoryClient

MAX_WORKERS = 10
MAX_RETRIES = 3


def _extract_custom_field(custom_fields: list[dict[str, Any]], api_name: str) -> str:
    for field in custom_fields or []:
        if field.get("api_name") == api_name:
            return str(field.get("value_formatted") or field.get("value") or "").strip()
    return ""


def fetch_item_detail(client: ZohoInventoryClient, organization_id: str, item_id: str) -> dict[str, Any]:
    last_exc: Exception | None = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            result = client._request_json("/items/" + item_id, {"organization_id": organization_id})
            return result.get("item", result)
        except Exception as exc:  # pragma: no cover - network flakiness
            last_exc = exc
            time.sleep(0.5 * attempt)
    raise RuntimeError(f"Failed to fetch item {item_id} after {MAX_RETRIES} attempts: {last_exc}")


def main() -> None:
    output_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[3] / "Zoho_Items_Export.xlsx"

    settings = get_settings()
    client = ZohoInventoryClient(
        client_id=settings.zoho_client_id,
        client_secret=settings.zoho_client_secret,
        refresh_token=settings.zoho_refresh_token,
        organization_id=settings.zoho_organization_id,
    )

    print("Fetching full item list from Zoho Inventory...")
    items = client.fetch_all_items()
    total = len(items)
    print(f"Found {total} items. Fetching per-item details for Product Category / Sub Category...")

    rows: list[dict[str, str]] = []
    errors: list[str] = []
    done = 0
    start_time = time.time()

    def process(item: dict[str, Any]) -> dict[str, str]:
        item_id = item.get("item_id") or ""
        detail = fetch_item_detail(client, settings.zoho_organization_id, item_id) if item_id else {}
        custom_fields = detail.get("custom_fields") or []
        return {
            "item_name": item.get("item_name") or item.get("name") or "",
            "sku": item.get("sku") or "",
            "product_category": _extract_custom_field(custom_fields, "cf_product_category"),
            "product_sub_category": _extract_custom_field(custom_fields, "cf_product_sub_category"),
            "show_in_online_store": "Yes" if item.get("show_in_storefront") is True else "No",
        }

    results_by_idx: dict[int, dict[str, str]] = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        futures = {pool.submit(process, item): idx for idx, item in enumerate(items)}
        for future in as_completed(futures):
            idx = futures[future]
            item = items[idx]
            done += 1
            try:
                results_by_idx[idx] = future.result()
            except Exception as exc:
                errors.append(f"{item.get('item_id')} ({item.get('sku')}): {exc}")
                results_by_idx[idx] = {
                    "item_name": item.get("item_name") or item.get("name") or "",
                    "sku": item.get("sku") or "",
                    "product_category": "ERROR",
                    "product_sub_category": "ERROR",
                    "show_in_online_store": "Yes" if item.get("show_in_storefront") is True else "No",
                }
            if done % 100 == 0 or done == total:
                elapsed = time.time() - start_time
                print(f"  {done}/{total} processed ({elapsed:.0f}s elapsed)")

    rows = [results_by_idx[idx] for idx in range(total)]

    if errors:
        print(f"Completed with {len(errors)} item(s) failing to fetch details:")
        for err in errors[:20]:
            print("  -", err)

    print("Writing Excel file...")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Zoho Items"

    headers = ["Item Name", "SKU", "Product Category", "Product Sub Category", "Show in Online Store"]
    sheet.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(
            [
                row["item_name"],
                row["sku"],
                row["product_category"],
                row["product_sub_category"],
                row["show_in_online_store"],
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    column_widths = [40, 28, 20, 22, 20]
    for idx, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"Saved {len(rows)} rows to {output_path}")


if __name__ == "__main__":
    main()
